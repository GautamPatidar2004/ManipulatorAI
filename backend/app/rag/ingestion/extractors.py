import io
import os
import re
import unicodedata
from typing import List, Optional, Tuple
from app.rag.ingestion.exceptions import (
    EmptyDocumentError,
    TextExtractionError,
    UnsupportedFileTypeError,
)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def sanitize_filename(filename: str) -> str:
    """
    Sanitize filename against directory traversal and dangerous characters.
    Extracts only the terminal basename and strips null bytes or traversal markers.
    """
    base = os.path.basename(filename.replace("\\", "/"))
    # Remove null bytes and control chars
    clean = re.sub(r"[\x00-\x1f\x7f]", "", base)
    clean = clean.strip()
    if not clean or clean in {".", ".."}:
        return "unnamed_document.txt"
    return clean


def normalize_text(text: str) -> str:
    """
    Clean and normalize extracted document text:
    - Normalizes unicode forms (NFKC).
    - Standardizes newline characters (\r\n -> \n).
    - Collapses multiple redundant blank lines.
    - Strips non-printable ASCII/control codes while preserving valid text.
    """
    if not text:
        return ""

    # Unicode normalization
    normalized = unicodedata.normalize("NFKC", text)

    # Standardize linebreaks
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")

    # Remove non-printable control characters except newline and tab
    normalized = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", normalized)

    # Collapse sequences of 3+ newlines to 2
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)

    return normalized.strip()


def extract_pdf_pages(file_bytes: bytes, filename: str) -> List[Tuple[Optional[int], str]]:
    """Extract page-indexed text passages from a PDF document using pypdf."""
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        reader = PdfReader(io.BytesIO(file_bytes))

        if reader.is_encrypted:
            try:
                # Try empty password decrypt if encrypted without pass
                if not reader.decrypt(""):
                    raise TextExtractionError(filename, "PDF is password-protected or encrypted.")
            except Exception:
                raise TextExtractionError(filename, "PDF is password-protected or encrypted.")

        pages_extracted: List[Tuple[Optional[int], str]] = []
        for idx, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            clean_page_text = normalize_text(page_text)
            if clean_page_text:
                pages_extracted.append((idx, clean_page_text))

        return pages_extracted

    except TextExtractionError:
        raise
    except Exception as exc:
        raise TextExtractionError(filename, f"PDF extraction failed: {str(exc)}")


def extract_docx_content(file_bytes: bytes, filename: str) -> List[Tuple[Optional[int], str]]:
    """Extract text passages from a DOCX file using python-docx."""
    try:
        import docx

        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs: List[str] = []

        # Extract regular paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)

        # Extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        full_text = "\n\n".join(paragraphs)
        clean = normalize_text(full_text)
        if not clean:
            return []

        return [(1, clean)]

    except Exception as exc:
        raise TextExtractionError(filename, f"DOCX extraction failed: {str(exc)}")


def extract_plain_or_markdown(file_bytes: bytes, filename: str) -> List[Tuple[Optional[int], str]]:
    """Extract text from TXT or Markdown files with fallback encoding detection."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "utf-16"]
    decoded: Optional[str] = None

    for enc in encodings:
        try:
            decoded = file_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if decoded is None:
        raise TextExtractionError(filename, "Failed to decode text file with supported encodings.")

    clean = normalize_text(decoded)
    if not clean:
        return []

    return [(1, clean)]


def extract_document(
    file_bytes: bytes,
    filename: str,
) -> List[Tuple[Optional[int], str]]:
    """
    Validate file extension and extract text content into page-indexed tuples.
    Raises UnsupportedFileTypeError if extension is invalid,
    TextExtractionError on parsing failures,
    and EmptyDocumentError if the document contains no text.
    """
    safe_name = sanitize_filename(filename)
    _, ext = os.path.splitext(safe_name)
    ext_lower = ext.lower()

    if ext_lower not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            filename=safe_name,
            extension=ext_lower,
            supported=sorted(list(SUPPORTED_EXTENSIONS)),
        )

    if ext_lower == ".pdf":
        pages = extract_pdf_pages(file_bytes, safe_name)
    elif ext_lower == ".docx":
        pages = extract_docx_content(file_bytes, safe_name)
    elif ext_lower in {".txt", ".md"}:
        pages = extract_plain_or_markdown(file_bytes, safe_name)
    else:
        raise UnsupportedFileTypeError(
            filename=safe_name,
            extension=ext_lower,
            supported=sorted(list(SUPPORTED_EXTENSIONS)),
        )

    # Check that text was actually extracted
    total_text = "".join(text for _, text in pages).strip()
    if not total_text:
        raise EmptyDocumentError(filename=safe_name)

    return pages
