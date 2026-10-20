import io
import os
from unittest.mock import AsyncMock, patch
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app.config import Settings
from app.db.base import Base
from app.main import app
from app.models.knowledge import KnowledgeChunk
from app.rag.embeddings.mock import MockEmbeddingProvider
from app.rag.ingestion.chunker import chunk_document_pages, split_text_into_chunks
from app.rag.ingestion.cli import run_cli_ingestion
from app.rag.ingestion.exceptions import (
    EmptyDocumentError,
    FileTooLargeError,
    IngestionAuthError,
    UnsupportedFileTypeError,
)
from app.rag.ingestion.extractors import extract_document, extract_docx_content, sanitize_filename
from app.rag.ingestion.service import IngestionService
from app.rag.vector_store import VectorStore

from sqlalchemy.pool import StaticPool

# In-memory SQLite engine for tests with shared static connection pool
test_engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)

Base.metadata.create_all(bind=test_engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


from app.db.session import get_db

app.dependency_overrides[get_db] = override_get_db
client = TestClient(app)


@pytest.fixture(autouse=True)
def clean_db():
    """Ensure tables exist and clean KnowledgeChunk table before each test run."""
    Base.metadata.create_all(bind=test_engine)
    app.dependency_overrides[get_db] = override_get_db
    db = TestingSessionLocal()
    try:
        db.query(KnowledgeChunk).delete()
        db.commit()
    finally:
        db.close()
    yield
    app.dependency_overrides[get_db] = override_get_db


# ---------------------------------------------------------------------------
# 1. Text Extraction Tests
# ---------------------------------------------------------------------------

def test_extract_markdown_and_plain_text():
    raw_md = b"# Engineering Career Track\n\nStaff engineers drive cross-team architecture.\n"
    pages = extract_document(raw_md, "career.md")
    assert len(pages) == 1
    assert pages[0][0] == 1
    assert "Staff engineers drive cross-team architecture." in pages[0][1]


def test_extract_docx_in_memory():
    import docx

    doc = docx.Document()
    doc.add_heading("Cloud Systems Guide", level=1)
    doc.add_paragraph("Reliability engineering is critical for modern cloud platforms.")

    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Metric"
    hdr_cells[1].text = "SLO Target"

    row_cells = table.add_row().cells
    row_cells[0].text = "Availability"
    row_cells[1].text = "99.99%"

    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()

    pages = extract_document(docx_bytes, "guide.docx")
    assert len(pages) == 1
    assert "Reliability engineering is critical" in pages[0][1]
    assert "Availability | 99.99%" in pages[0][1]


def test_extract_pdf_pages_in_memory():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    pdf_buf = io.BytesIO()
    writer.write(pdf_buf)
    pdf_bytes = pdf_buf.getvalue()

    # Empty blank page raises EmptyDocumentError
    with pytest.raises(EmptyDocumentError):
        extract_document(pdf_bytes, "blank.pdf")


def test_sanitize_filename_traversal():
    assert sanitize_filename("../../../etc/passwd") == "passwd"
    assert sanitize_filename("..\\..\\windows\\system32\\cmd.exe") == "cmd.exe"
    assert sanitize_filename("safe_file.pdf") == "safe_file.pdf"


# ---------------------------------------------------------------------------
# 2. Chunking & Overlap Tests
# ---------------------------------------------------------------------------

def test_chunking_and_overlap_boundaries():
    text = (
        "Paragraph 1 discusses software architecture fundamentals.\n\n"
        "Paragraph 2 discusses distributed systems and consensus protocols.\n\n"
        "Paragraph 3 discusses observability and automated alerting."
    )
    chunks = split_text_into_chunks(text, page_number=1, chunk_size=80, chunk_overlap=20)
    assert len(chunks) >= 2
    for c in chunks:
        assert c.page_number == 1
        assert len(c.chunk_text) <= 90  # Respects chunk size limit allowance
        assert c.chunk_text != ""

    # Verify global indexing preservation
    doc_pages = [(1, "Page one text"), (2, "Page two text")]
    all_chunks = chunk_document_pages(doc_pages, chunk_size=50, chunk_overlap=10)
    assert len(all_chunks) == 2
    assert all_chunks[0].chunk_index == 0
    assert all_chunks[0].page_number == 1
    assert all_chunks[1].chunk_index == 1
    assert all_chunks[1].page_number == 2


# ---------------------------------------------------------------------------
# 3. Idempotency & Duplicate Detection Tests
# ---------------------------------------------------------------------------

@pytest.mark.anyio
async def test_duplicate_detection_idempotency():
    db = TestingSessionLocal()
    try:
        sample_text = b"Idempotent document test text for knowledge indexing."
        cfg = Settings(EMBEDDING_PROVIDER="mock", EMBEDDING_DIM=384)

        # First ingestion
        res1 = await IngestionService.ingest_document(
            db=db,
            file_bytes=sample_text,
            filename="test_doc.txt",
            source_id="test_idempotency_doc",
            settings=cfg,
        )
        assert res1.status == "created"
        assert res1.chunks_created >= 1

        chunks_after_first = VectorStore.get_chunks_by_source(db, "test_idempotency_doc")
        assert len(chunks_after_first) == res1.chunks_created

        # Second ingestion with exact same content
        res2 = await IngestionService.ingest_document(
            db=db,
            file_bytes=sample_text,
            filename="test_doc.txt",
            source_id="test_idempotency_doc",
            settings=cfg,
        )
        assert res2.status == "skipped_unchanged"
        assert res2.chunks_created == 0
        assert res2.total_chunks == res1.chunks_created

        # Verify no duplicate rows were added
        chunks_after_second = VectorStore.get_chunks_by_source(db, "test_idempotency_doc")
        assert len(chunks_after_second) == len(chunks_after_first)

    finally:
        db.close()


# ---------------------------------------------------------------------------
# 4. Changed Documents Purge Stale Chunks Test
# ---------------------------------------------------------------------------

@pytest.mark.anyio
async def test_changed_document_purges_stale_chunks():
    db = TestingSessionLocal()
    try:
        cfg = Settings(
            EMBEDDING_PROVIDER="mock",
            EMBEDDING_DIM=384,
            INGESTION_CHUNK_SIZE=50,
            INGESTION_CHUNK_OVERLAP=10,
        )

        # Version 1: Long text that generates multiple chunks
        v1_text = (
            b"First long chapter on executive engineering strategy.\n\n"
            b"Second long chapter on career milestones and leadership.\n\n"
            b"Third long chapter on organizational scaling and influence."
        )
        res1 = await IngestionService.ingest_document(
            db=db,
            file_bytes=v1_text,
            filename="versioned.txt",
            source_id="versioned_doc",
            settings=cfg,
        )
        assert res1.status == "created"
        v1_count = res1.chunks_created
        assert v1_count >= 2

        # Version 2: Modified, short text under the same source_id (fits in 1 chunk)
        v2_text = b"Updated single short document."
        res2 = await IngestionService.ingest_document(
            db=db,
            file_bytes=v2_text,
            filename="versioned.txt",
            source_id="versioned_doc",
            settings=cfg,
        )
        assert res2.status == "updated"

        # Check that old chunks were purged and only new chunks exist
        chunks_after_update = VectorStore.get_chunks_by_source(db, "versioned_doc")
        assert len(chunks_after_update) == res2.chunks_created
        assert len(chunks_after_update) == 1
        assert "Updated single short document." in chunks_after_update[0].chunk_text

    finally:
        db.close()


# ---------------------------------------------------------------------------
# 5. Metadata Preservation Test
# ---------------------------------------------------------------------------

@pytest.mark.anyio
async def test_metadata_preservation_in_knowledge_chunks():
    db = TestingSessionLocal()
    try:
        sample_content = b"# Executive Guide\n\nPrincipal engineers lead technical innovation."
        cfg = Settings(EMBEDDING_PROVIDER="mock", EMBEDDING_DIM=384)

        result = await IngestionService.ingest_document(
            db=db,
            file_bytes=sample_content,
            filename="exec_guide.md",
            source_id="exec_guide_2026",
            source_title="Executive Guide 2026",
            source_url="https://example.com/exec-guide",
            extra_metadata={"track": "executive", "department": "engineering"},
            settings=cfg,
        )

        chunks = VectorStore.get_chunks_by_source(db, "exec_guide_2026")
        assert len(chunks) >= 1
        meta = chunks[0].source_metadata

        assert meta["doc_hash"] == result.doc_hash
        assert meta["file_name"] == "exec_guide.md"
        assert meta["file_type"] == "md"
        assert meta["chunk_index"] == 0
        assert meta["track"] == "executive"
        assert meta["department"] == "engineering"
        assert "ingested_at" in meta
        assert chunks[0].source_title == "Executive Guide 2026"
        assert chunks[0].source_url == "https://example.com/exec-guide"

    finally:
        db.close()


# ---------------------------------------------------------------------------
# 6. Failure Modes & Validation Tests
# ---------------------------------------------------------------------------

@pytest.mark.anyio
async def test_unsupported_file_extension():
    db = TestingSessionLocal()
    try:
        with pytest.raises(UnsupportedFileTypeError) as exc_info:
            await IngestionService.ingest_document(
                db=db,
                file_bytes=b"executable content",
                filename="script.exe",
            )
        assert exc_info.value.code == "UNSUPPORTED_FILE_TYPE"
        assert exc_info.value.status_code == 400
    finally:
        db.close()


@pytest.mark.anyio
async def test_empty_file_rejected():
    db = TestingSessionLocal()
    try:
        with pytest.raises(EmptyDocumentError) as exc_info:
            await IngestionService.ingest_document(
                db=db,
                file_bytes=b"   \n\n\t  ",
                filename="empty.txt",
            )
        assert exc_info.value.code == "EMPTY_DOCUMENT"
        assert exc_info.value.status_code == 400
    finally:
        db.close()


@pytest.mark.anyio
async def test_file_size_limit_exceeded():
    db = TestingSessionLocal()
    try:
        cfg = Settings(MAX_UPLOAD_SIZE_BYTES=100)
        large_bytes = b"A" * 150

        with pytest.raises(FileTooLargeError) as exc_info:
            await IngestionService.ingest_document(
                db=db,
                file_bytes=large_bytes,
                filename="large.txt",
                settings=cfg,
            )
        assert exc_info.value.code == "FILE_TOO_LARGE"
        assert exc_info.value.status_code == 413
    finally:
        db.close()


# ---------------------------------------------------------------------------
# 7. Ingestion REST API Tests
# ---------------------------------------------------------------------------

def test_api_upload_and_source_retrieval():
    fixture_path = os.path.join(os.path.dirname(__file__), "fixtures", "synthetic_sample.md")
    with open(fixture_path, "rb") as f:
        file_bytes = f.read()

    # Upload document
    resp = client.post(
        "/api/v1/ingestion/upload",
        files={"file": ("synthetic_sample.md", file_bytes, "text/markdown")},
        data={"source_id": "synthetic_delta9", "source_title": "Synthetic Delta 9 Guide"},
    )
    assert resp.status_code == 201
    data = resp.json()
    assert data["source_id"] == "synthetic_delta9"
    assert data["status"] == "created"
    assert data["chunks_created"] >= 1

    # Inspect source metadata
    detail_resp = client.get("/api/v1/ingestion/sources/synthetic_delta9")
    assert detail_resp.status_code == 200
    detail = detail_resp.json()
    assert detail["source_id"] == "synthetic_delta9"
    assert detail["chunk_count"] == data["chunks_created"]

    # Delete source
    del_resp = client.delete("/api/v1/ingestion/sources/synthetic_delta9")
    assert del_resp.status_code == 200
    assert del_resp.json()["deleted_chunks"] == data["chunks_created"]


def test_api_upload_with_auth_key():
    from app.config import get_settings

    secure_settings = Settings(
        INGESTION_API_KEY="secret-ingest-key-123",
        EMBEDDING_PROVIDER="mock",
        EMBEDDING_DIM=384,
    )
    app.dependency_overrides[get_settings] = lambda: secure_settings

    try:
        # 1. Missing header -> 403 Forbidden
        bad_resp = client.post(
            "/api/v1/ingestion/upload",
            files={"file": ("test.txt", b"sample content", "text/plain")},
        )
        assert bad_resp.status_code == 403

        # 2. Correct header -> 201 Created
        good_resp = client.post(
            "/api/v1/ingestion/upload",
            headers={"X-Ingestion-Key": "secret-ingest-key-123"},
            files={"file": ("test.txt", b"sample content", "text/plain")},
        )
        assert good_resp.status_code == 201

    finally:
        app.dependency_overrides.pop(get_settings, None)


# ---------------------------------------------------------------------------
# 8. CLI Ingestion Test
# ---------------------------------------------------------------------------

@pytest.mark.anyio
async def test_cli_runner_ingests_synthetic_fixture():
    fixture_path = os.path.join(os.path.dirname(__file__), "fixtures", "synthetic_sample.md")
    with patch("app.rag.ingestion.cli.get_session_factory", return_value=TestingSessionLocal):
        exit_code = await run_cli_ingestion(
            target_path=fixture_path,
            source_id="cli_fixture_test",
            source_title="CLI Fixture Test",
        )
        assert exit_code == 0

        # Running a second time should succeed and report skipped
        exit_code_second = await run_cli_ingestion(
            target_path=fixture_path,
            source_id="cli_fixture_test",
        )
        assert exit_code_second == 0


@pytest.mark.anyio
async def test_embedding_failure_rolls_back_cleanly():
    """Verify that an embedding provider error raises exception and leaves zero chunks."""
    db = TestingSessionLocal()
    try:
        from app.rag.exceptions import EmbeddingProviderError

        with patch("app.rag.ingestion.service.get_embedding_provider") as mock_get_provider:
            mock_provider = AsyncMock()
            mock_provider.embed_batch.side_effect = EmbeddingProviderError("embedding API timeout")
            mock_get_provider.return_value = mock_provider

            with pytest.raises(EmbeddingProviderError):
                await IngestionService.ingest_document(
                    db=db,
                    file_bytes=b"Some document text",
                    filename="fail.txt",
                    source_id="fail_source_id",
                )

        # Zero chunks should be persisted
        assert len(VectorStore.get_chunks_by_source(db, "fail_source_id")) == 0

    finally:
        db.close()

